from __future__ import annotations

import io
import json
import logging
import os
import re
import tempfile
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import fitz
import psutil
import pytesseract
from huggingface_hub import hf_hub_download
from PIL import Image
from docx import Document
from llama_cpp import Llama
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

logger = logging.getLogger(__name__)

OCR_LANG = os.getenv("EMISSION_OCR_LANG", "rus+eng")
GGUF_REPO = os.getenv("EMISSION_GGUF_REPO", "Qwen/Qwen3-0.6B-GGUF")
GGUF_FILE = os.getenv("EMISSION_GGUF_FILE", "Qwen3-0.6B-Q4_K_M.gguf")
GGUF_REVISION = os.getenv("EMISSION_GGUF_REVISION", "1208e45d782fe18602c5eaf10e5758d5b0f24c03")
LOCAL_MAX_NEW_TOKENS = int(os.getenv("EMISSION_LOCAL_MAX_NEW_TOKENS", "700"))
LOCAL_CONTEXT_TOKENS = int(os.getenv("EMISSION_LOCAL_CONTEXT_TOKENS", "4096"))
MIN_FREE_RAM_MB = int(os.getenv("EMISSION_MIN_FREE_RAM_MB", "850"))
ANALYSIS_MAX_CHARS = int(os.getenv("EMISSION_ANALYSIS_MAX_CHARS", "9000"))


@dataclass
class SourcePage:
    page: int
    text: str
    method: str
    confidence: float | None = None


@dataclass
class DocumentData:
    filename: str
    file_type: str
    pages: list[SourcePage] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(
            f"[Страница {p.page}]\n{p.text}" for p in self.pages if p.text.strip()
        )


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _preprocess_ocr_image(image: Image.Image) -> Image.Image:
    gray = image.convert("L")
    gray = gray.resize((int(gray.width * 1.15), int(gray.height * 1.15)))
    return gray.point(lambda p: 255 if p > 210 else (0 if p < 115 else p))


def _ocr_page(page: fitz.Page) -> tuple[str, float | None]:
    pix = page.get_pixmap(matrix=fitz.Matrix(2.6, 2.6), alpha=False)
    image = _preprocess_ocr_image(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))
    candidates: list[tuple[str, float | None]] = []
    for psm in (6, 4, 11):
        try:
            data = pytesseract.image_to_data(
                image, lang=OCR_LANG, config=f"--oem 1 --psm {psm}",
                output_type=pytesseract.Output.DICT, timeout=60,
            )
        except (pytesseract.TesseractError, RuntimeError) as exc:
            logger.warning("Tesseract failed psm=%s: %s", psm, exc)
            continue
        parts, confs = [], []
        for token, conf in zip(data.get("text", []), data.get("conf", [])):
            token = str(token).strip()
            if not token:
                continue
            parts.append(token)
            try:
                value = float(conf)
                if value >= 0:
                    confs.append(value)
            except (TypeError, ValueError):
                pass
        text = _clean_text(" ".join(parts))
        if text:
            candidates.append((text, round(sum(confs) / len(confs), 1) if confs else None))
    if not candidates:
        return "", None
    return max(candidates, key=lambda x: (len(re.sub(r"\s+", "", x[0])), x[1] or 0))


def extract_pdf(data: bytes, filename: str) -> DocumentData:
    doc = fitz.open(stream=data, filetype="pdf")
    pages: list[SourcePage] = []
    try:
        for idx, page in enumerate(doc, 1):
            text = _clean_text(page.get_text("text"))
            if len(re.sub(r"\s+", "", text)) < 80:
                text, confidence = _ocr_page(page)
                method = "OCR" if text else "empty"
            else:
                confidence, method = None, "text-layer"
            pages.append(SourcePage(idx, text, method, confidence))
    finally:
        doc.close()
    return DocumentData(filename, "PDF", pages)


def extract_docx(data: bytes, filename: str) -> DocumentData:
    with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
        tmp.write(data)
        tmp.flush()
        doc = Document(tmp.name)
        blocks: list[str] = []
        for paragraph in doc.paragraphs:
            value = _clean_text(paragraph.text)
            if value:
                blocks.append(value)
        for table_index, table in enumerate(doc.tables, 1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(_clean_text(cell.text) for cell in row.cells))
            if rows:
                blocks.append(f"[Таблица {table_index}]\n" + "\n".join(rows))
    return DocumentData(filename, "DOCX", [SourcePage(1, "\n\n".join(blocks), "docx")])


def extract_document(data: bytes, filename: str) -> DocumentData:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(data, filename)
    if suffix == ".docx":
        return extract_docx(data, filename)
    raise ValueError("Поддерживаются PDF и DOCX. Формат DOC пока не поддерживается.")


RELEVANT_TERMS = (
    "эмитент", "выпуск", "объем", "номинал", "количество", "размещен",
    "размещение", "цена", "оплата", "погашен", "погашение", "срок обращения",
    "купон", "ставка", "амортиза", "оферт", "досроч", "выкуп", "дефолт",
    "ковенант", "обеспеч", "гарант", "субординац", "очередност", "налог",
    "комисси", "расход", "ограничен", "риск", "платеж", "дата", "формул",
    "уведомлен", "требован", "досрочн", "обязательств", "право", "условия"
)


def _keyword_fragments(document: DocumentData, max_chars: int = ANALYSIS_MAX_CHARS,
                       window_chars: int = 850) -> tuple[str, list[dict[str, Any]]]:
    hits: list[tuple[int, int, str, list[str]]] = []
    for page in document.pages:
        text = re.sub(r"\s+", " ", page.text or "").strip()
        if not text:
            continue
        lower = text.lower()
        positions: list[tuple[int, str]] = []
        for term in RELEVANT_TERMS:
            start = 0
            while True:
                pos = lower.find(term, start)
                if pos < 0:
                    break
                positions.append((pos, term))
                start = pos + len(term)
        positions.sort()
        ranges: list[tuple[int, int, list[str]]] = []
        for pos, term in positions:
            left, right = max(0, pos - window_chars), min(len(text), pos + len(term) + window_chars)
            if ranges and left <= ranges[-1][1] + 250:
                a, b, terms = ranges[-1]
                ranges[-1] = (a, max(b, right), terms + [term])
            else:
                ranges.append((left, right, [term]))
        for left, right, terms in ranges:
            fragment = text[left:right].strip()
            score = len(set(terms)) + (3 if page.page == 1 else 0)
            hits.append((score, page.page, fragment, sorted(set(terms))))
    hits.sort(key=lambda x: (x[0], -x[1]), reverse=True)
    selected, meta, total, seen = [], [], 0, set()
    for score, page_no, fragment, terms in hits:
        key = (page_no, fragment[:180])
        if key in seen:
            continue
        block = f"[Страница {page_no}]\n{fragment}"
        if total + len(block) > max_chars:
            continue
        seen.add(key)
        selected.append(block)
        meta.append({"page": page_no, "keywords": terms, "score": score, "fragment": fragment})
        total += len(block)
    return "\n\n".join(sorted(selected, key=lambda x: int(re.search(r"\d+", x).group()))), sorted(meta, key=lambda x: x["page"])


def find_keyword_fragments(document: DocumentData, max_chars: int = ANALYSIS_MAX_CHARS):
    return _keyword_fragments(document, max_chars=max_chars)


COMBINED_SYSTEM_PROMPT = """
Ты анализируешь один эмиссионный документ. Используй только переданный текст.
Не используй рынок, MOEX, новости или внешние сведения.

Извлеки существенные условия выпуска и затем сделай краткий нейтральный анализ.
Для каждого факта сохраняй страницу. Сложные условия НЕ сокращай: перечисляй
все указанные триггеры, даты, сроки, формулы, порядок уведомления, цены,
суммы, ограничения и процедуры. Если данных нет, напиши "не указано в документе".

Верни только JSON следующего вида:
{
"facts":{"issuer":"...","security_type":"...","issue_volume":"...","nominal":"...","currency":"...","maturity":"...","placement_period":"...","coupon":"...","coupon_periods":"...","amortization":"...","offers":"...","early_redemption":"...","buyback":"...","security_guarantees":"...","subordination":"...","covenants":"...","default_events":"...","payment_mechanics":"...","claim_priority":"...","taxes_legal_notes":"...","fees_expenses":"...","restrictions":"...","risk_factors":"...","unusual_terms":"...","other_material_terms":"...","source_references":[{"page":1,"topic":"..."}]},
"analysis":{"summary":"...","key_features":[],"attention_points":[],"risks":[],"holder_favorable_mechanisms":[],"holder_unfavorable_terms":[],"analytical_conclusion":"...","detailed_conditions":{"early_redemption":"...","offers":"...","default_events":"...","covenants":"...","amortization":"..."}}
}
""".strip()


_QWEN: Llama | None = None
_QWEN_LOAD_LOCK = threading.Lock()


def _get_qwen(progress=None) -> Llama:
    global _QWEN
    if _QWEN is not None:
        return _QWEN

    # Streamlit can execute multiple sessions/reruns in the same process.
    # Never let two requests download/load the model simultaneously.
    with _QWEN_LOAD_LOCK:
        if _QWEN is not None:
            return _QWEN

        available_mb = psutil.virtual_memory().available / (1024 * 1024)
        if available_mb < MIN_FREE_RAM_MB:
            raise RuntimeError(
                f"Недостаточно свободной RAM для Qwen3 GGUF: доступно {available_mb:.0f} МБ, "
                f"требуется минимум {MIN_FREE_RAM_MB} МБ. "
                "Распознавание текста и поиск по ключевым словам доступны без AI."
            )

        if progress:
            progress("Загружаем файл Qwen3-0.6B Q4 GGUF…")
        try:
            model_path = hf_hub_download(
                repo_id=GGUF_REPO,
                filename=GGUF_FILE,
                revision=GGUF_REVISION,
                etag_timeout=10,
            )
        except Exception as exc:
            raise RuntimeError(
                f"Не удалось скачать Qwen3 GGUF: {exc}. "
                "Проверьте доступ к Hugging Face и имя GGUF-файла."
            ) from exc

        if not Path(model_path).is_file() or Path(model_path).stat().st_size < 100 * 1024 * 1024:
            raise RuntimeError(f"GGUF-файл скачан некорректно или слишком мал: {model_path}")

        if progress:
            progress("Файл Qwen3 загружен. Инициализируем llama.cpp…")
        try:
            model = Llama(
                model_path=model_path,
                n_ctx=LOCAL_CONTEXT_TOKENS,
                n_threads=max(1, min(2, os.cpu_count() or 2)),
                n_batch=32,
                verbose=False,
            )
        except Exception as exc:
            raise RuntimeError(f"Не удалось инициализировать llama.cpp: {exc}") from exc

        _QWEN = model
        return _QWEN

def _empty_analysis(keyword_hits):
    missing = "не указано в документе"
    keys = ("issuer","security_type","issue_volume","nominal","currency","maturity",
            "placement_period","coupon","coupon_periods","amortization","offers",
            "early_redemption","buyback","security_guarantees","subordination",
            "covenants","default_events","payment_mechanics","claim_priority",
            "taxes_legal_notes","fees_expenses","restrictions","risk_factors",
            "unusual_terms","other_material_terms")
    return {
        "keyword_hits": keyword_hits,
        "facts": {k: missing for k in keys} | {"source_references": []},
        "analysis": {
            "summary": missing, "key_features": [], "attention_points": [], "risks": [],
            "holder_favorable_mechanisms": [], "holder_unfavorable_terms": [],
            "analytical_conclusion": missing,
            "detailed_conditions": {k: missing for k in (
                "early_redemption","offers","default_events","covenants","amortization"
            )}
        }
    }

def _extract_json_object(raw):
    text = (raw or "").strip()
    try:
        value = json.loads(text)
        if isinstance(value, dict):
            return value
    except json.JSONDecodeError:
        pass
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        raise RuntimeError("Qwen3 не вернул JSON-объект.")
    try:
        value = json.loads(text[start:end + 1])
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Qwen3 вернул некорректный JSON: {exc}") from exc
    if not isinstance(value, dict):
        raise RuntimeError("Qwen3 вернул JSON не в виде объекта.")
    return value

def _normalise_analysis(value, keyword_hits):
    result = _empty_analysis(keyword_hits)
    if isinstance(value.get("facts"), dict):
        result["facts"].update(value["facts"])
    if isinstance(value.get("analysis"), dict):
        result["analysis"].update(value["analysis"])
        detailed = value["analysis"].get("detailed_conditions")
        if isinstance(detailed, dict):
            result["analysis"]["detailed_conditions"].update(detailed)
    for key, val in list(result["facts"].items()):
        if key == "source_references":
            if not isinstance(val, list):
                result["facts"][key] = []
        elif val is None or (isinstance(val, str) and not val.strip()):
            result["facts"][key] = "не указано в документе"
    for key in ("key_features","attention_points","risks",
                "holder_favorable_mechanisms","holder_unfavorable_terms"):
        if not isinstance(result["analysis"].get(key), list):
            result["analysis"][key] = []
    if not result["analysis"].get("summary"):
        result["analysis"]["summary"] = "не указано в документе"
    if not result["analysis"].get("analytical_conclusion"):
        result["analysis"]["analytical_conclusion"] = "не указано в документе"
    detailed = result["analysis"].setdefault("detailed_conditions", {})
    for key in ("early_redemption","offers","default_events","covenants","amortization"):
        if not detailed.get(key):
            detailed[key] = result["facts"].get(key, "не указано в документе")
    return result

def _llm_complete(system_prompt, user_prompt, progress=None):
    if progress:
        progress("Шаг 2/2: запускаем Qwen3 через llama.cpp…")
    llm = _get_qwen(progress=progress)
    try:
        response = llm.create_chat_completion(
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt + "\n\n/no_think\nВерни только JSON."},
            ],
            max_tokens=LOCAL_MAX_NEW_TOKENS,
            temperature=0.0,
            top_p=0.8,
            response_format={"type": "json_object"},
        )
    except Exception as exc:
        raise RuntimeError(f"Ошибка генерации Qwen3: {exc}") from exc
    choices = response.get("choices") or []
    if not choices:
        raise RuntimeError("Qwen3 не вернул результат.")
    content = (choices[0].get("message") or {}).get("content")
    if not isinstance(content, str) or not content.strip():
        raise RuntimeError("Qwen3 вернул пустой ответ.")
    if progress:
        progress("Qwen3 завершил анализ.")
    return content.strip()

def analyze_document(document, progress=None):
    if document is None or not document.pages:
        raise ValueError("Документ не загружен.")
    if progress:
        progress("Шаг 1/2: ищем важные фрагменты по ключевым словам…")
    fragments, keyword_hits = _keyword_fragments(
        document, max_chars=ANALYSIS_MAX_CHARS, window_chars=650
    )
    if not fragments:
        return _empty_analysis(keyword_hits)
    if progress:
        progress(f"Шаг 1/2 завершён: найдено {len(keyword_hits)} фрагментов. Передаём контекст Qwen3…")
    prompt = (
        "Анализируй только фрагменты одного документа ниже. Страницы указаны явно. "
        "Не используй внешние сведения. Не сокращай материальные условия. "
        "Для досрочного погашения, оферт, дефолтов, ковенант и амортизации "
        "перечисляй все доступные триггеры, даты, сроки, формулы, цены, суммы, "
        "порядок заявления, уведомления и исполнения. Не заменяй это общей фразой. "
        "Если данных нет, пиши «не указано в документе».\n\nФРАГМЕНТЫ:\n" + fragments
    )
    raw = _llm_complete(COMBINED_SYSTEM_PROMPT, prompt, progress)
    return _normalise_analysis(_extract_json_object(raw), keyword_hits)


def _pdf_p(text, style):
    value = str(text if text not in (None, "") else "не указано в документе")
    value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(value.replace("\n", "<br/>"), style)

def build_analysis_report(document, result):
    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "EmissionTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=16, leading=20, alignment=TA_CENTER, spaceAfter=12
    )
    h2 = ParagraphStyle(
        "EmissionH2", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=11, leading=14, spaceBefore=8, spaceAfter=4
    )
    body = ParagraphStyle(
        "EmissionBody", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.5, leading=11, spaceAfter=4
    )

    pdf = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=12*mm, leftMargin=12*mm,
        topMargin=12*mm, bottomMargin=12*mm,
        title="Анализ эмиссионного документа"
    )
    story = [
        _pdf_p("Анализ эмиссионного документа", title),
        _pdf_p("Документ: " + document.filename, body),
        _pdf_p(
            "Источник: только загруженный документ. Рыночные данные и MOEX не используются.",
            body
        ),
    ]
    facts = result.get("facts", {})
    analysis = result.get("analysis", {})

    story.append(_pdf_p("1. Краткое описание", h2))
    story.append(_pdf_p(analysis.get("summary"), body))

    story.append(_pdf_p("2. Основные параметры", h2))
    labels = [
        ("Эмитент", "issuer"), ("Вид бумаги", "security_type"),
        ("Объём выпуска", "issue_volume"), ("Номинал", "nominal"),
        ("Валюта", "currency"), ("Погашение", "maturity"),
        ("Размещение", "placement_period"), ("Купон", "coupon"),
        ("Периоды купона", "coupon_periods"), ("Амортизация", "amortization"),
        ("Оферты", "offers"), ("Досрочное погашение", "early_redemption"),
        ("Выкуп", "buyback"), ("Обеспечение/гарантии", "security_guarantees"),
        ("Субординация", "subordination"), ("Ковенанты", "covenants"),
        ("События дефолта", "default_events"), ("Платёжный механизм", "payment_mechanics"),
        ("Очередность требований", "claim_priority"), ("Налоги/правовые примечания", "taxes_legal_notes"),
        ("Комиссии/расходы", "fees_expenses"), ("Ограничения", "restrictions"),
        ("Риски из документа", "risk_factors"), ("Прочие существенные условия", "other_material_terms"),
    ]
    rows = [[_pdf_p("Показатель", body), _pdf_p("Условие", body)]]
    for label, key in labels:
        rows.append([_pdf_p(label, body), _pdf_p(facts.get(key), body)])
    table = Table(rows, colWidths=[48*mm, 136*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0,0), (-1,-1), 0.3, colors.grey),
        ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 3),
        ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("TOPPADDING", (0,0), (-1,-1), 3),
        ("BOTTOMPADDING", (0,0), (-1,-1), 3),
    ]))
    story.append(table)

    sections = [
        ("3. Ключевые особенности", "key_features"),
        ("4. На что обратить внимание", "attention_points"),
        ("5. Риски", "risks"),
        ("6. Потенциально неблагоприятные условия для держателя", "holder_unfavorable_terms"),
        ("7. Механизмы защиты / потенциально благоприятные условия", "holder_favorable_mechanisms"),
    ]
    for heading, key in sections:
        story.append(_pdf_p(heading, h2))
        values = analysis.get(key) or []
        if isinstance(values, list):
            for item in values or ["не указано в документе"]:
                story.append(_pdf_p("• " + str(item), body))
        else:
            story.append(_pdf_p(values, body))

    story.append(_pdf_p("8. Подробные существенные условия", h2))
    detailed = analysis.get("detailed_conditions") or {}
    for label, key in (
        ("Досрочное погашение", "early_redemption"),
        ("Оферты", "offers"),
        ("События дефолта", "default_events"),
        ("Ковенанты", "covenants"),
        ("Амортизация", "amortization"),
    ):
        story.append(_pdf_p(label, h2))
        story.append(_pdf_p(detailed.get(key) or facts.get(key), body))

    story.append(_pdf_p("9. Аналитическое заключение", h2))
    story.append(_pdf_p(analysis.get("analytical_conclusion"), body))

    story.append(_pdf_p("10. Источники по страницам", h2))
    refs = facts.get("source_references") or []
    if isinstance(refs, list) and refs:
        for ref in refs:
            if isinstance(ref, dict):
                story.append(_pdf_p(
                    "Страница " + str(ref.get("page", "не указана")) +
                    ": " + str(ref.get("topic", "не указано")),
                    body
                ))
    else:
        for page in document.pages:
            if page.text.strip():
                story.append(_pdf_p("Страница " + str(page.page), body))

    pdf.build(story)
    return buffer.getvalue()
