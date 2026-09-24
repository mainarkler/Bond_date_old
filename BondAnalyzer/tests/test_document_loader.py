from pathlib import Path
from docx import Document
from app.core.document_loader import load_document

def test_docx_paragraphs_and_tables(tmp_path):
 path=tmp_path/'bond.docx'; doc=Document(); doc.add_paragraph('Эмитент: Тест'); table=doc.add_table(rows=1, cols=2); table.cell(0,0).text='Номинал'; table.cell(0,1).text='1000 RUB'; doc.save(path)
 loaded=load_document(path); assert loaded.pages[0].page==1 and '1000 RUB' in loaded.pages[0].text

def test_empty_docx(tmp_path):
 path=tmp_path/'empty.docx'; Document().save(path); assert load_document(path).pages[0].text==''

def test_pdf_preserves_page_numbers(tmp_path):
 import fitz
 path=tmp_path/'bond.pdf'; pdf=fitz.open(); first=pdf.new_page(); first.insert_text((72,72),'First page'); second=pdf.new_page(); second.insert_text((72,72),'Second page'); pdf.save(path); pdf.close()
 loaded=load_document(path); assert [page.page for page in loaded.pages]==[1,2] and loaded.pages[1].text=='Second page'
