from pathlib import Path
from docx import Document
from .models import LoadedDocument, PageText

def load_docx(path: str | Path) -> LoadedDocument:
    path = Path(path); document = Document(path)
    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells): parts.append(" | ".join(cells))
    return LoadedDocument(path.name, [PageText(1, "\n".join(parts))], {"format": "docx", "logical_pages": True})
